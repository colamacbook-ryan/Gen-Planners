#!/usr/bin/env python3
"""Gen-Planners CLI — plan-as-code.

SSOT của mỗi dự án nằm trong projects/<slug>/:
  assumptions.yaml  mọi con số (giả định) + model tài chính + kịch bản
  evidence.yaml     nguồn/bằng chứng mà assumptions trích dẫn
  sections/*.md     văn bản hồ sơ (Jinja) — chỉ tham chiếu số, không gõ số
  deck.md           pitch deck (Jinja, pandoc slides)
build/ là output sinh ra — không bao giờ sửa tay.

Lệnh:
  python gen.py new <slug> --name "Tên dự án" [--idea-file f.md]
  python gen.py build <slug>     # tính model, render hồ sơ + deck + excel
  python gen.py check <slug>     # kiểm tra SSOT, trả exit 1 nếu có lỗi
  python gen.py status [<slug>]  # tiến độ pipeline
"""
import argparse
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

import yaml
from jinja2 import Environment, FileSystemLoader, StrictUndefined

ROOT = Path(__file__).resolve().parent
PROJECTS = Path(os.environ.get("GP_PROJECTS") or ROOT / "projects")  # SaaS trỏ workspace riêng
TEMPLATE = ROOT / "templates" / "project"
STAGES = yaml.safe_load((ROOT / "pipeline" / "stages.yaml").read_text())["stages"]


def rel(p: Path) -> str:
    return str(p.relative_to(ROOT)) if p.is_relative_to(ROOT) else str(p)


def load(p: Path):
    return yaml.safe_load(p.read_text()) or {}


def project_dir(slug: str) -> Path:
    d = PROJECTS / slug
    if not d.exists():
        sys.exit(f"Không tìm thấy dự án: {d}")
    return d


# ---------------------------------------------------------------- model

def per_year(v, n):
    """Scalar -> lặp n năm; list -> giữ nguyên (pad bằng phần tử cuối)."""
    if isinstance(v, list):
        return (v + [v[-1]] * n)[:n]
    return [v] * n


def compute(assumptions: dict, scenario: dict) -> dict:
    meta, A, M = assumptions["meta"], assumptions["assumptions"], assumptions["model"]
    n, start = meta["years"], meta["start_year"]

    def val(ref):
        """Chuỗi theo năm từ 1 ref: số | list số | id assumption | "cum:<id>" (luỹ kế)
        | list id (nhân từng phần tử, vd giá × tỷ lệ)."""
        if isinstance(ref, list) and ref and all(isinstance(r, str) for r in ref):
            out = [1.0] * n
            for r in ref:
                out = [a * b for a, b in zip(out, val(r))]
            return out
        if isinstance(ref, str) and ref.startswith("cum:"):
            out, acc = [], 0.0
            for x in val(ref[4:]):
                acc += x
                out.append(acc)
            return out
        return per_year(A[ref]["value"] if isinstance(ref, str) else ref, n)

    years = [dict(year=start + i, revenue=0.0, cogs=0.0, opex=0.0, streams={}, revenue_one_off=0.0, gp_one_off=0.0)
             for i in range(n)]
    segs = {}
    seg = lambda sid: segs.setdefault(sid, dict(
        id=sid, name=M.get("segments", {}).get(sid, sid),
        years=[dict(year=start + i, revenue=0.0, cogs=0.0, direct_opex=0.0) for i in range(n)]))
    stream_x = scenario.get("stream_x", {})
    for s in M["revenue_streams"]:
        price, vol, cogs = val(s["price"]), val(s["volume"]), val(s["cogs_pct"])
        k = scenario.get("volume_x", 1) * scenario.get("price_x", 1) * stream_x.get(s["id"], 1)
        sy = seg(s.get("segment", s["id"]))["years"]
        for i, y in enumerate(years):
            r = price[i] * vol[i] * k
            y["streams"][s["id"]] = r
            y["revenue"] += r
            y["cogs"] += r * cogs[i]
            if s.get("kind") == "one_off":           # bán một lần (thiết bị, phí gia nhập…) — không lặp lại, tách khỏi định giá
                y["revenue_one_off"] += r
                y["gp_one_off"] += r * (1 - cogs[i])
            sy[i]["revenue"] += r
            sy[i]["cogs"] += r * cogs[i]
    for o in M["opex"]:
        for i, x in enumerate(val(o["amount"])):
            x *= scenario.get("opex_x", 1)
            years[i]["opex"] += x
            if o.get("segment"):
                seg(o["segment"])["years"][i]["direct_opex"] += x
    for g in segs.values():
        for sy in g["years"]:
            sy["gross_profit"] = sy["revenue"] - sy["cogs"]
            sy["contribution"] = sy["gross_profit"] - sy["direct_opex"]
            sy["margin"] = sy["contribution"] / sy["revenue"] if sy["revenue"] else 0
        g["last"] = g["years"][-1]
        g["total_revenue"] = sum(sy["revenue"] for sy in g["years"])
        g["total_contribution"] = sum(sy["contribution"] for sy in g["years"])

    tax_rate = val(M["tax_rate"])
    capex = val(M["capex"])
    raise_amt = A[M["funding"]]["value"]
    founder = A[M["founder_capital"]]["value"] if M.get("founder_capital") else 0   # vốn founder tự góp (tuỳ chọn)
    cash, loss_cf, cum_fcf = raise_amt + founder, 0.0, -raise_amt - founder
    for i, y in enumerate(years):
        y["gross_profit"] = y["revenue"] - y["cogs"]
        y["gross_margin"] = y["gross_profit"] / y["revenue"] if y["revenue"] else 0
        y["ebitda"] = y["gross_profit"] - y["opex"]
        y["revenue_recurring"] = y["revenue"] - y["revenue_one_off"]
        y["recurring_share"] = y["revenue_recurring"] / y["revenue"] if y["revenue"] else 0
        y["ebitda_recurring"] = y["ebitda"] - y["gp_one_off"]     # EBITDA nếu bỏ lãi gộp từ khoản bán một lần
        taxable = y["ebitda"] + loss_cf          # chuyển lỗ đơn giản
        y["tax"] = max(taxable, 0) * tax_rate[i]
        loss_cf = min(taxable, 0)
        y["net_income"] = y["ebitda"] - y["tax"]
        y["capex"] = capex[i]
        y["fcf"] = y["net_income"] - y["capex"]
        cash += y["fcf"]
        cum_fcf += y["fcf"]
        y["cash_end"] = cash
        y["cum_fcf"] = cum_fcf

    pre = A[M["pre_money"]]["value"]
    stake = raise_amt / (pre + raise_amt)
    basis = M.get("exit_basis", "ebitda")        # recurring_ebitda: chỉ định giá phần lặp lại (thận trọng, chuẩn khi có doanh thu một lần)
    exit_ebitda = years[-1]["ebitda_recurring" if basis == "recurring_ebitda" else "ebitda"]
    exit_value = max(exit_ebitda, 0) * A[M["exit_multiple"]]["value"]
    r = A[M["discount_rate"]]["value"]
    first = lambda cond: next((y["year"] for y in years if cond(y)), None)
    metrics = dict(
        raise_amount=raise_amt,
        founder_capital=founder,
        pre_money=pre,
        post_money=pre + raise_amt,
        investor_stake=stake,
        breakeven_year=first(lambda y: y["ebitda"] > 0),
        payback_year=first(lambda y: y["cum_fcf"] >= 0),
        min_cash=min(y["cash_end"] for y in years),
        funding_gap=max(0, -min(y["cash_end"] for y in years)),
        npv=-raise_amt + sum(y["fcf"] / (1 + r) ** (i + 1) for i, y in enumerate(years)),
        exit_basis=basis,
        exit_ebitda=exit_ebitda,
        recurring_share=years[-1]["recurring_share"],
        one_off_revenue_total=sum(y["revenue_one_off"] for y in years),
        exit_value=exit_value,
        investor_exit=exit_value * stake,
        investor_moic=exit_value * stake / raise_amt if raise_amt else 0,
    )
    return dict(years=years, last=years[-1], metrics=metrics, segments=segs)


def compute_all(assumptions: dict) -> dict:
    scen = assumptions["model"].get("scenarios", {"base": {}})
    return {k: compute(assumptions, v or {}) for k, v in scen.items()}


# ---------------------------------------------------------------- render

def make_env(d: Path, assumptions: dict, evidence: dict, issues: list):
    unit = assumptions["meta"].get("unit_label", "")

    def auto(x, digits):
        """Không chỉ định số lẻ: số lớn hoặc số tròn thì bỏ phần lẻ (1.500 thay vì 1.500,0)."""
        if digits is not None:
            return digits
        return 0 if abs(x) >= 100 or float(x).is_integer() else 1

    def money(x, digits=None):
        return f"{x:,.{auto(x, digits)}f} {unit}".replace(",", "_").replace(".", ",").replace("_", ".").strip()

    def num(x, digits=0):
        return f"{x:,.{digits}f}".replace(",", "_").replace(".", ",").replace("_", ".")

    def ty(x, digits=None):
        """triệu đ → tỷ đ."""
        return num(x / 1000, auto(x / 1000, digits)) + " tỷ đ"

    def pct(x, digits=0):
        return f"{x * 100:.{digits}f}%".replace(".", ",")

    def cite(*ids):
        """Trích kiểu tác giả–năm cho người đọc: (Momentum Works, 2025; JLL, 2026). Danh mục đầy đủ có URL ở phụ lục/Excel."""
        out = []
        for i in ids:
            e = evidence.get(i)
            if e is None:
                issues.append(f"cite('{i}') không có trong evidence.yaml")
                out.append(i)
                continue
            who = str(e.get("publisher") or e.get("title") or i).split(" — ")[0].strip()
            who = who if len(who) <= 40 else who[:38].rstrip() + "…"
            year = str(e.get("date") or "")[:4]
            out.append(f"{who}, {year}" if year.isdigit() else who)
        return "(" + "; ".join(out) + ")"

    env = Environment(loader=FileSystemLoader(str(d)), undefined=StrictUndefined,
                      trim_blocks=True, lstrip_blocks=True)
    env.filters.update(money=money, num=num, pct=pct, ty=ty)
    env.globals.update(cite=cite)
    return env


def context(d: Path):
    assumptions, evidence = load(d / "assumptions.yaml"), load(d / "evidence.yaml")
    project = load(d / "project.yaml")
    A = assumptions["assumptions"]
    unverified = [k for k, v in A.items() if v.get("status") != "verified"]
    ctx = dict(
        p=project,
        meta=assumptions["meta"],
        a={k: v["value"] for k, v in A.items()},
        A=A,
        ev=evidence,
        fin=compute_all(assumptions),
        unverified=unverified,
        draft=bool(unverified),
        today=dt.date.today().isoformat(),
        version=f'{project.get("version", "0.1")}+{git_rev(d if (d / ".git").exists() else ROOT)}',
    )
    return ctx, assumptions, evidence


def git_rev(where=ROOT):
    """Mã phiên bản: commit của repo dự án (SaaS) hoặc của repo chuẩn."""
    try:
        return subprocess.check_output(["git", "-c", "safe.directory=*", "rev-parse", "--short", "HEAD"], cwd=where,
                                       stderr=subprocess.DEVNULL, text=True).strip()
    except Exception:
        return "nogit"


def cmd_build(slug):
    d = project_dir(slug)
    ctx, assumptions, evidence = context(d)
    issues = []
    env = make_env(d, assumptions, evidence, issues)
    out = d / "build"
    shutil.rmtree(out, ignore_errors=True)
    out.mkdir()

    sections = sorted((d / "sections").glob("*.md"))
    plan = "\n\n".join(env.get_template(f"sections/{s.name}").render(**ctx) for s in sections)
    (out / "business-plan.md").write_text(plan)
    (out / "deck.md").write_text(env.get_template("deck.md").render(**ctx))
    write_excel(out / "financials.xlsx", ctx, assumptions, evidence)

    pandoc = shutil.which("pandoc")
    if pandoc:
        subprocess.run([pandoc, "business-plan.md", "-o", "business-plan.docx", "--toc"], cwd=out, check=True)
        subprocess.run([pandoc, "deck.md", "-o", "deck.pptx", "--slide-level=2"], cwd=out, check=True)
    else:
        print("! Không có pandoc — chỉ xuất .md/.xlsx")

    for i in issues:
        print("! " + i)
    print(f"✓ build {slug} v{ctx['version']} → {rel(out)}/")
    for f in sorted(out.iterdir()):
        print("   ", f.name)
    if ctx["draft"]:
        print(f"! DRAFT: {len(ctx['unverified'])} giả định chưa verified — chưa dùng để gọi vốn.")


def write_excel(path, ctx, assumptions, evidence):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Assumptions"
    ws.append(["id", "label", "value", "unit", "status", "confidence", "source", "note"])
    for k, v in assumptions["assumptions"].items():
        val = v["value"] if not isinstance(v["value"], list) else ", ".join(map(str, v["value"]))
        ws.append([k, v.get("label"), val, v.get("unit"), v.get("status"), v.get("confidence"),
                   ", ".join(v.get("source", []) or []), v.get("note")])
    rows = [("Doanh thu", "revenue"), ("  trong đó định kỳ", "revenue_recurring"), ("  trong đó bán một lần", "revenue_one_off"),
            ("Giá vốn", "cogs"), ("Lợi nhuận gộp", "gross_profit"),
            ("Chi phí vận hành", "opex"), ("EBITDA", "ebitda"),
            ("EBITDA định kỳ (bỏ lãi gộp bán một lần)", "ebitda_recurring"), ("Thuế", "tax"),
            ("Lợi nhuận ròng", "net_income"), ("Capex", "capex"), ("Free cash flow", "fcf"),
            ("Tiền cuối kỳ", "cash_end")]
    for name, f in ctx["fin"].items():
        s = wb.create_sheet(f"P&L {name}")
        s.append(["Chỉ tiêu"] + [y["year"] for y in f["years"]])
        for label, key in rows:
            s.append([label] + [round(y[key], 2) for y in f["years"]])
        s.append([])
        for k, v in f["metrics"].items():
            s.append([k, v])
        for c in s[1]:
            c.font = Font(bold=True)
    s = wb.create_sheet("Segments base")
    base = ctx["fin"]["base"]
    s.append(["Mảng", "Chỉ tiêu"] + [y["year"] for y in base["years"]])
    for g in base["segments"].values():
        for label, key in [("Doanh thu", "revenue"), ("Lợi nhuận gộp", "gross_profit"),
                           ("Chi phí trực tiếp", "direct_opex"), ("Lợi nhuận đóng góp", "contribution")]:
            s.append([g["name"], label] + [round(y[key], 2) for y in g["years"]])
    e = wb.create_sheet("Evidence")
    e.append(["id", "title", "publisher", "date", "url", "type", "note"])
    for k, v in evidence.items():
        e.append([k, v.get("title"), v.get("publisher"), str(v.get("date", "")), v.get("url"),
                  v.get("type"), v.get("note")])
    wb.save(path)


# ---------------------------------------------------------------- check

HARDCODED = re.compile(r"\d[\d.,]*\s*(%|tỷ|triệu|nghìn|USD|VND|VNĐ|đ\b|\$|k\b|M\b)", re.I)


# Cấu trúc dữ liệu mỗi loại biểu đồ dashboard (khớp studio/index.html). Trường → kiểu phần tử bắt buộc.
CHART_SCHEMA = {
    "bar": {"items": ("label", "value")}, "dots": {"items": ("label", "value")}, "funnel": {"items": ("label", "value")},
    "donut": {"items": ("label", "value")}, "range": {"items": ("label", "lo", "hi")},
    "line": {"series": ("name", "points")}, "stacked": {"categories": (), "series": ("name", "values")},
    "grouped": {"categories": (), "series": ("name", "values")}, "waterfall": {"steps": ("label", "value")},
}
KPI_KEYS = {"label", "value", "note", "status", "source", "unit", "sources"}


def check_datasets(d):
    """datasets.yaml do agent viết: sai cấu trúc thì dashboard không vẽ được — báo lỗi cụ thể thay vì để trình xem hỏng."""
    errs = []
    ds = load(d / "datasets.yaml") if (d / "datasets.yaml").exists() else {}
    for i, k in enumerate(ds.get("kpis") or []):
        extra = set(k) - KPI_KEYS if isinstance(k, dict) else set()
        if extra:
            errs.append(f"datasets.yaml kpis[{i}]: khoá lạ {sorted(extra)} — thường do dấu phẩy trong giá trị, "
                        f"hãy đặt trong ngoặc kép, ví dụ value: \"66,8 USD/m²\"")
    for c in ds.get("charts") or []:
        cid, t = c.get("id", "?"), c.get("type")
        if t not in CHART_SCHEMA:
            errs.append(f"datasets.yaml chart {cid}: type '{t}' không hỗ trợ — dùng {', '.join(CHART_SCHEMA)}")
            continue
        for field, need in CHART_SCHEMA[t].items():
            v = c.get(field)
            if not isinstance(v, list) or not v:
                errs.append(f"datasets.yaml chart {cid} ({t}): cần '{field}' là danh sách"
                            + (f" các {{{', '.join(need)}}}" if need else ""))
                continue
            for j, it in enumerate(v):
                miss = [x for x in need if not (isinstance(it, dict) and x in it)]
                if miss:
                    errs.append(f"datasets.yaml chart {cid}: {field}[{j}] thiếu {miss}")
                    break
        if t in ("stacked", "grouped") and isinstance(c.get("series"), list) and isinstance(c.get("categories"), list):
            for sr in c["series"]:
                if isinstance(sr, dict) and len(sr.get("values") or []) != len(c["categories"]):
                    errs.append(f"datasets.yaml chart {cid}: series '{sr.get('name')}' cần {len(c['categories'])} giá trị theo categories")
    return errs


INTERNAL_EVIDENCE = {"internal", "founder", "interview", "estimate"}


STREAM_KINDS = ("recurring", "one_off")


def sanity(fin):
    """Cảnh báo số liệu khó tin — nhà đầu tư sẽ hỏi ngay; không chặn build, chỉ nhắc kiểm lại giả định."""
    out, base = [], fin.get("base")
    if not base:
        return out
    Y, M = base["years"], base["metrics"]
    last = Y[-1]
    if last["revenue"] and last["ebitda"] / last["revenue"] > 0.45:
        out.append(f"hợp lý: biên EBITDA năm {last['year']} = {last['ebitda'] / last['revenue']:.0%} — rất cao so với phần lớn ngành, kiểm lại giá vốn/chi phí")
    if Y[0]["ebitda"] > 0 and Y[0]["revenue"]:
        out.append(f"hợp lý: có lãi EBITDA ngay năm đầu {Y[0]['year']} — hiếm gặp với dự án mới, kiểm lại chi phí khởi động")
    for a, b in zip(Y[1:], Y[2:]):
        if a["revenue"] > 0 and b["revenue"] / a["revenue"] > 4:
            out.append(f"hợp lý: doanh thu {b['year']} gấp {b['revenue'] / a['revenue']:.1f} lần {a['year']} — tăng trưởng rất nóng, cần lý do cụ thể")
    one_off = 1 - last["recurring_share"] if last["revenue"] else 0
    if one_off > 0.2 and M.get("exit_basis") == "ebitda":
        out.append(f"hợp lý: {one_off:.0%} doanh thu năm {last['year']} là bán một lần nhưng định giá thoát vốn vẫn nhân trên toàn bộ EBITDA "
                   f"— đặt model.exit_basis: recurring_ebitda (EBITDA định kỳ năm cuối = {last['ebitda_recurring']:,.0f})".replace(",", "."))
    if M.get("investor_moic", 0) > 20:
        out.append(f"hợp lý: MOIC nhà đầu tư {M['investor_moic']:.0f}× — quá đẹp để tin, kiểm lại sản lượng và hệ số thoát vốn")
    if M.get("payback_year") and M["payback_year"] <= Y[0]["year"]:
        out.append("hợp lý: hoàn vốn ngay trong năm đầu — kiểm lại")
    for name, f in fin.items():
        if name != "base" and f["metrics"]["funding_gap"] > 0:
            out.append(f"kịch bản {name}: thiếu tiền {f['metrics']['funding_gap']:,.0f} — nêu rõ trong mục rủi ro và phương án dự phòng".replace(",", "."))
    return out


def strings(x, path="", parent=None):
    """Duyệt mọi chuỗi trong YAML (story/datasets/products) kèm đường dẫn và dict chứa nó."""
    if isinstance(x, dict):
        for k, v in x.items():
            yield from strings(v, f"{path}.{k}" if path else str(k), x)
    elif isinstance(x, list):
        for i, v in enumerate(x):
            yield from strings(v, f"{path}[{i}]", parent)
    elif isinstance(x, str):
        yield path, x, parent or {}


SOURCED = ("src", "source", "sources", "url")


def cmd_check(slug, as_json=False):
    d = project_dir(slug)
    errors, warns = [], []
    assumptions, evidence = load(d / "assumptions.yaml"), load(d / "evidence.yaml")
    errors += check_datasets(d)
    ds = load(d / "datasets.yaml") if (d / "datasets.yaml").exists() else {}
    for i, k in enumerate(ds.get("kpis") or []):
        for sid in (k.get("source") or []) if isinstance(k, dict) and isinstance(k.get("source"), list) else []:
            if sid not in evidence:
                errors.append(f"datasets.yaml kpis[{i}]: source '{sid}' không có trong evidence.yaml")

    for k, v in assumptions["assumptions"].items():
        src = v.get("source") or []
        for s in src:
            if s not in evidence:
                errors.append(f"assumption {k}: source '{s}' không tồn tại trong evidence.yaml")
        if v.get("status") == "verified" and not src:
            errors.append(f"assumption {k}: verified nhưng không có source")
        elif v.get("status") == "verified" and all(s == "ev-000" or (evidence.get(s) or {}).get("type") in INTERNAL_EVIDENCE for s in src):
            errors.append(f"assumption {k}: verified nhưng chỉ dựa trên nguồn nội bộ/lời founder ({', '.join(src)}) — để 'estimate' tới khi có nguồn ngoài")
        if v.get("status") in (None, "placeholder"):
            warns.append(f"assumption {k}: PLACEHOLDER — cần nghiên cứu")
    for k, v in evidence.items():
        if not v.get("url") and v.get("type") not in ("interview", "quote", "internal"):
            warns.append(f"evidence {k}: thiếu url")

    for f in sorted((d / "sections").glob("*.md")) + [d / "deck.md"]:
        for n, line in enumerate(f.read_text().splitlines(), 1):
            bare = re.sub(r"\{\{.*?\}\}|\{%.*?%\}|\{#.*?#\}|\[ev-[\w-]+\]", "", line)
            if HARDCODED.search(bare):
                errors.append(f"{f.relative_to(d)}:{n}: số liệu gõ tay → đưa vào assumptions.yaml: {line.strip()[:80]}")
            if "TODO" in line:
                warns.append(f"{f.relative_to(d)}:{n}: TODO")

    try:
        ctx, _, _ = context(d)
        issues = []
        env = make_env(d, assumptions, evidence, issues)
        for s in sorted((d / "sections").glob("*.md")):
            env.get_template(f"sections/{s.name}").render(**ctx)
        env.get_template("deck.md").render(**ctx)
        errors += issues
        M = assumptions.get("model", {})
        for st in M.get("revenue_streams", []):
            if st.get("kind", "recurring") not in STREAM_KINDS:
                errors.append(f"model: revenue_streams {st.get('id')}: kind '{st.get('kind')}' không hợp lệ — dùng {' | '.join(STREAM_KINDS)}")
        if M.get("exit_basis", "ebitda") not in ("ebitda", "recurring_ebitda"):
            errors.append(f"model: exit_basis '{M.get('exit_basis')}' không hợp lệ — dùng ebitda | recurring_ebitda")
        base = ctx["fin"].get("base")
        if base and base["metrics"]["funding_gap"] > 0:
            errors.append(f"model: kịch bản base âm tiền {base['metrics']['funding_gap']:,.0f} — số vốn gọi không đủ")
        warns += sanity(ctx["fin"])
        # story/datasets/products: chuỗi có {{ … }} được render theo model khi xuất — lỗi phải lộ ra ở đây
        hand = []
        for name in ("story.yaml", "datasets.yaml", "products.yaml"):
            if not (d / name).exists():
                continue
            for path, text, parent in strings(load(d / name)):
                if "{{" in text or "{%" in text:
                    try:
                        env.from_string(text).render(**ctx)
                    except Exception as e:  # noqa: BLE001
                        errors.append(f"{name} {path}: biểu thức Jinja lỗi: {e}")
                elif (name != "products.yaml" and re.search(r"(takeaway|headline|thesis|sub|value)$", path) and HARDCODED.search(text)
                      and parent.get("status") != "actual" and not any(parent.get(k) for k in SOURCED)):   # dữ kiện thực tế có nguồn: được gõ
                    hand.append(f"{name} {path}")
        if hand:
            warns.append(f"{len(hand)} chỗ trong story/datasets gõ tay số tiền/% (ví dụ {', '.join(hand[:3])}) — nếu là kết quả model, "
                         "dùng {{ fin.base.last.revenue|ty }}… để tự cập nhật; nếu là dữ kiện có nguồn thì giữ nguyên")
    except Exception as e:  # noqa: BLE001
        errors.append(f"render/model lỗi: {type(e).__name__}: {e}")

    if as_json:
        print(json.dumps({"errors": errors, "warnings": warns}, ensure_ascii=False))
        return 1 if errors else 0
    for w in warns:
        print("  warn ", w)
    for e in errors:
        print("  ERROR", e)
    print(f"{'✗' if errors else '✓'} check {slug}: {len(errors)} lỗi, {len(warns)} cảnh báo")
    return 1 if errors else 0


# ---------------------------------------------------------------- new / status

def cmd_new(slug, name, idea_file):
    d = PROJECTS / slug
    if d.exists():
        sys.exit(f"Đã tồn tại: {d}")
    shutil.copytree(TEMPLATE, d)
    p = load(d / "project.yaml")
    p.update(name=name, slug=slug, created=dt.date.today().isoformat(),
             stages={s["id"]: "todo" for s in STAGES})
    (d / "project.yaml").write_text(yaml.safe_dump(p, allow_unicode=True, sort_keys=False))
    if idea_file:
        (d / "idea.md").write_text(Path(idea_file).read_text())
    print(f"✓ tạo {rel(d)} — bước tiếp: điền idea.md rồi chạy stage 01")


def cmd_status(slug):
    slugs = [slug] if slug else sorted(p.name for p in PROJECTS.iterdir() if (p / "project.yaml").exists())
    for s in slugs:
        d = project_dir(s)
        p = load(d / "project.yaml")
        A = load(d / "assumptions.yaml")["assumptions"]
        by = {}
        for v in A.values():
            by[v.get("status", "placeholder")] = by.get(v.get("status", "placeholder"), 0) + 1
        done = sum(1 for v in p["stages"].values() if v == "done")
        print(f"\n{p['name']} ({s}) — stage {done}/{len(STAGES)} · assumptions {by}")
        for st in STAGES:
            mark = {"done": "✓", "doing": "…", "todo": " "}.get(p["stages"].get(st["id"]), "?")
            print(f"  [{mark}] {st['id']}  {st['name']}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    n = sub.add_parser("new")
    n.add_argument("slug")
    n.add_argument("--name", required=True)
    n.add_argument("--idea-file")
    sub.add_parser("build").add_argument("slug")
    ck = sub.add_parser("check")
    ck.add_argument("slug")
    ck.add_argument("--json", action="store_true", help="in lỗi/cảnh báo dạng JSON (cho SaaS)")
    sub.add_parser("status").add_argument("slug", nargs="?")
    a = ap.parse_args()
    if a.cmd == "new":
        cmd_new(a.slug, a.name, a.idea_file)
    elif a.cmd == "build":
        cmd_build(a.slug)
    elif a.cmd == "check":
        sys.exit(cmd_check(a.slug, a.json))
    else:
        cmd_status(a.slug)


if __name__ == "__main__":
    main()
